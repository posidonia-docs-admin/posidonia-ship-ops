# -*- coding: utf-8 -*-
"""Le a planilha "Analise Viagens - Historico.xlsx" e a traduz para o modelo do Corsair.

So LEITURA e validacao — nada aqui toca o banco. Quem grava e o
importar_historico.py, e so depois de o relatorio desta leitura ser aprovado.

O LAYOUT (uma aba por navio, um diario vertical):

    C=porto   D=VOY      E=ACTION    F=data       G=hora     H=data+hora  I=TEMPO   K..=notas
              APN24001   SAILING     23/11/2024   15:00

Blocos por parada, cada um aberto por uma linha de cabecalho (E = "ACTION"). O
porto do bloco esta na coluna C — na propria linha de cabecalho, ou numa linha
abaixo (varia por aba). O Courage tem uma coluna vazia a mais: a data esta em G.
A coluna A carrega ORION/PASSAGEM/LOADING/DISCHARGING e a tonelagem (com "mt"
na B); a linha "ROB" e o carregado menos o descarregado.

As paradas de bunker (Mosqueiro/Icoaraci, Itaqui) aparecem DENTRO do bloco de
saida de Alumar, depois do Sailing, em quatro formatos: FUNDEADO / INICIO /
TERMINO / SAILING; ARRIVAL / BERTH / SAILING; duas linhas sem rotulo (inicio e
termino do abastecimento); ou INICIO / TERMINO seguidos de uma linha sem rotulo
(a saida). As trocas de tripulacao na descida aparecem no bloco de Juruti,
depois do Sailing, como um par de linhas sem rotulo.

AS SEIS DECISOES (Vinicius, 14/09/2026):

  1. Unberth nunca foi registrado (Juruti e Alumar): COPIA-SE a hora do Sailing,
     com observacao no marco. A duracao da viagem passa a ser Sailing -> Sailing,
     como as supervisoras ja calculam.
  2. Fazendinha na descida e Barra Norte nao tem registro: as escalas nascem
     CANCELADAS com a observacao "sem registro historico". Onde houve registro
     (troca de tripulacao na descida, Barra Norte na subida), entram.
  3. Todo marco importado nasce CONFERIDO pela conta "importacao".
  4. O que existe no sistema e teste: a importacao apaga tudo antes de gravar.
  5. Commander e Courage comecam no Rio de Janeiro: a escala de abertura da
     primeira viagem e cancelada com a nota da saida do Rio.
  6. Erros de digitacao conhecidos sao corrigidos em AJUSTES, abaixo, e nao na
     planilha — para o original continuar sendo o original.
"""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field

from app import dominio

OFFSET = "-03:00"
DATA_DECISOES = "14/09/2026"

# aba (normalizada) -> navio_id do seed
NAVIOS = {
    "AMAZON PATHFINDER": 1,
    "AMAZON PIONEER": 2,
    "AMAZON COMMANDER": 3,
    "AMAZON COURAGE": 4,
}
PREFIXOS = {1: "APT", 2: "APN", 3: "ACM", 4: "ACR"}

# Portos que abrem bloco. A chave e o comeco do texto normalizado da coluna C.
PORTOS_PRINCIPAIS = (
    ("ALUMAR", "ALUMAR"),
    ("MACAPA", "FAZENDINHA"),
    ("FAZENDINHA", "FAZENDINHA"),
    ("JURUTI", "JURUTI"),
    ("RIO DE JANEIRO", "RIO"),
)
# Portos das paradas adicionais, dentro do bloco de saida de Alumar.
# None = "fundeadouro" sem dizer qual: presume-se pelas horas desde a saida.
PORTOS_EXTRA = (
    ("MOSQUEIRO", "ICOARACI"),
    ("ICOARACI", "ICOARACI"),
    ("ITAQUI", "ITAQUI"),
    ("BARRA NORTE", "BARRA_NORTE"),
    ("FUNDEADOR", None),
)
# Abaixo disto desde a saida de Alumar, um fundeadouro sem nome e Itaqui (ao
# lado); acima, e a regiao de Mosqueiro/Icoaraci (~30 h de navegacao).
HORAS_ITAQUI = 20

CHAVES_MODELO = ("abertura", "faz_subida", "juruti", "faz_descida", "barra_norte", "encerramento")
ORDEM_MODELO = {"abertura": 10, "faz_subida": 20, "juruti": 30,
                "faz_descida": 40, "barra_norte": 50, "encerramento": 60}

RE_CODIGO = re.compile(r"^A[A-Z]{2}\d{5}$")
RE_BUNKER = re.compile(r"ABAST|ABST|BUNKER|VLSFO|MGO", re.I)
RE_DATA = re.compile(r"\d{1,2}/\d{1,2}(?:/\d{2,4})?|\d{1,2}:\d{2}")
RE_NUMERO = re.compile(r"\d+(?:[.,]\d+)?")

# ---------------------------------------------------------------------------
# Ajustes: erros de digitacao confirmados pelo Vinicius. (aba, linha, coluna,
# valor que a celula tem hoje, valor certo). Se a celula ja nao tiver o valor
# antigo (alguem corrigiu a planilha), o ajuste e pulado e o relatorio avisa.
# ---------------------------------------------------------------------------
AJUSTES = (
    ("AMAZON PATHFINDER", 68, "F", dt.datetime(2024, 1, 3), dt.datetime(2025, 1, 3),
     "APT24003, Sailing de Juruti: ano digitado como 2024"),
    # As duas datas de Fazendinha estavam trocadas (Vinicius, 14/09/2026).
    ("AMAZON PIONEER", 120, "F", dt.datetime(2025, 2, 4), dt.datetime(2025, 2, 3),
     "APN25002, Arrival em Fazendinha: 03/02 e nao 04/02"),
    ("AMAZON PIONEER", 123, "F", dt.datetime(2025, 2, 3), dt.datetime(2025, 2, 4),
     "APN25002, Sailing de Fazendinha: 04/02 e nao 03/02"),
)

# Marcos que a planilha registra fora de qualquer celula que o leitor entenda.
# (aba, codigo, chave da escala, tipo, hora local ISO, motivo)
MARCOS_MANUAIS = (
    ("AMAZON COURAGE", "ACR26003", "extra:BARRA_NORTE", "sailing", "2026-02-09T18:00",
     "Saida de Barra Norte anotada na linha de cabecalho (H529), sem data: mesma data da chegada"),
)


# ---------------------------------------------------------------------------
# O modelo lido
# ---------------------------------------------------------------------------

@dataclass
class Marco:
    quando: dt.datetime            # hora local
    linha: int
    precisao: str = "exata"
    rotulo: str = ""               # o que a planilha chamou esta linha
    nota: str | None = None        # vira evento.observacao
    rob_vlsfo: float | None = None # so o CSV plano traz estes tres
    rob_mgo: float | None = None
    fw: float | None = None

    @property
    def iso(self) -> str:
        return self.quando.strftime("%Y-%m-%dT%H:%M")


@dataclass
class Escala:
    chave: str                     # chave do modelo, ou "extra"
    linha: int = 0
    porto: str | None = None       # so nas extras
    motivo: str = "bunker"         # so nas extras
    marcos: dict = field(default_factory=dict)      # tipo -> Marco
    notas: list = field(default_factory=list)
    abastecimento: tuple | None = None              # (vlsfo, mgo, texto)
    tonelagem: float | None = None
    cancelar: str | None = None    # motivo do cancelamento, se for o caso
    porto_presumido: bool = False

    @property
    def rotulo(self) -> str:
        if self.chave == "extra":
            return "parada adicional em {}".format(self.porto or "?")
        return {"abertura": "saída de Alumar", "faz_subida": "Fazendinha (subida)",
                "juruti": "Juruti", "faz_descida": "Fazendinha (descida)",
                "barra_norte": "Barra Norte", "encerramento": "Alumar (descarga)"}[self.chave]


@dataclass
class Viagem:
    navio_id: int
    codigo: str
    aba: str
    linha: int
    escalas: dict = field(default_factory=dict)     # chave -> Escala
    extras: list = field(default_factory=list)
    rob_planilha: float | None = None
    saida_registrada: Marco | None = None           # o SAILING do bloco de descarga
    aberta: bool = False

    def todos_os_marcos(self):
        """(escala, tipo, marco) na ordem da viagem, so os que existem."""
        ordem = ["abertura"] + ["extra"] + ["faz_subida", "juruti", "faz_descida",
                                             "barra_norte", "encerramento"]
        for chave in ordem:
            escalas = self.extras if chave == "extra" else [self.escalas[chave]]
            for escala in escalas:
                if escala.cancelar:
                    continue
                for tipo in ("arrival", "berth", "unberth", "sailing"):
                    if tipo in escala.marcos:
                        yield escala, tipo, escala.marcos[tipo]


@dataclass
class Problema:
    nivel: str                     # erro | aviso | nota
    aba: str
    codigo: str | None
    linha: int | None
    texto: str

    def __str__(self) -> str:
        onde = " ".join(p for p in (self.aba.replace("AMAZON ", "").title(),
                                    self.codigo or "", "L{}".format(self.linha) if self.linha else "") if p)
        return "{:<28} {}".format(onde, self.texto)


# ---------------------------------------------------------------------------
# Leitura celula a celula
# ---------------------------------------------------------------------------

@dataclass
class _Linha:
    n: int
    codigo: str | None
    acao: str | None               # arrival|berth|unberth|sailing|fundeado|inicio|termino|cabecalho|None
    quando: dt.datetime | None
    precisao: str
    notas: list
    porto_principal: str | None
    porto_extra: str | None        # nome do porto extra ou "" quando e fundeadouro sem nome
    tem_porto_extra: bool
    tonelagem: float | None
    rob: float | None
    combinada: dt.datetime | None  # a coluna data+hora, quando discorda


def _norm(texto) -> str:
    return re.sub(r"\s+", " ", dominio.normalizar(str(texto)).replace("\n", " ")).strip()


def _acao(valor) -> str | None:
    if not isinstance(valor, str):
        return None
    t = _norm(valor)
    if t in ("ARRIVAL", "BERTH", "UNBERTH", "SAILING"):
        return t.lower()
    if t == "ACTION":
        return "cabecalho"
    if t.startswith("FUNDEADO") or t.startswith("POSICAO"):
        return "fundeado"
    if t.startswith("INICIO"):
        return "inicio"
    if t.startswith("TERMINO") or t.startswith("FIM"):
        return "termino"
    return None


def _porto(valor) -> tuple[str | None, str | None, bool]:
    """(porto principal, porto extra, tem porto extra) a partir da coluna C."""
    if not isinstance(valor, str) or not valor.strip():
        return None, None, False
    t = _norm(valor)
    for comeco, codigo in PORTOS_PRINCIPAIS:
        if t.startswith(comeco):
            return codigo, None, False
    for comeco, codigo in PORTOS_EXTRA:
        if t.startswith(comeco):
            return None, codigo, True
    return None, None, False


def _quando(data, hora) -> tuple[dt.datetime | None, str]:
    if not isinstance(data, dt.datetime):
        return None, "exata"
    base = dt.datetime(data.year, data.month, data.day)
    if isinstance(hora, dt.time):
        return base.replace(hour=hora.hour, minute=hora.minute), "exata"
    if isinstance(hora, dt.timedelta):          # "24:00" vira 1 dia
        return base + hora, "exata"
    if isinstance(hora, dt.datetime):
        return base.replace(hour=hora.hour, minute=hora.minute), "exata"
    return base, "apenas_data"


def _nota(valor) -> str | None:
    if isinstance(valor, dt.datetime):
        return valor.strftime("%d/%m/%Y %H:%M" if (valor.hour or valor.minute) else "%d/%m/%Y")
    if isinstance(valor, str):
        t = re.sub(r"\s+", " ", valor).strip(" ,")
        if t and t.upper() != "TEMPO":
            return t
    return None


def _aplicar_ajustes(ws, aba: str, problemas: list) -> None:
    from openpyxl.utils import column_index_from_string
    for aba_alvo, linha, coluna, antigo, novo, motivo in AJUSTES:
        if aba_alvo != aba:
            continue
        celula = ws.cell(linha, column_index_from_string(coluna))
        if celula.value == antigo:
            celula.value = novo
            problemas.append(Problema("nota", aba, None, linha,
                                      "ajuste aplicado: {} ({} -> {})".format(
                                          motivo, _nota(antigo), _nota(novo))))
        else:
            problemas.append(Problema("aviso", aba, None, linha,
                                      "ajuste NAO aplicado, a celula {}{} tem {!r} e nao {!r}: {}".format(
                                          coluna, linha, celula.value, antigo, motivo)))


def _ler_linhas(ws, aba: str, problemas: list) -> list[_Linha]:
    col_data = None
    for r in range(1, min(ws.max_row, 80) + 1):
        if _acao(ws.cell(r, 5).value) == "sailing":
            for c in (6, 7):
                if isinstance(ws.cell(r, c).value, dt.datetime):
                    col_data = c
                    break
        if col_data:
            break
    if col_data is None:
        raise ValueError("Aba {}: nao achei a coluna de data.".format(aba))

    linhas = []
    for r in range(1, ws.max_row + 1):
        a, b, c, d, e = (ws.cell(r, i).value for i in range(1, 6))
        data, hora = ws.cell(r, col_data).value, ws.cell(r, col_data + 1).value
        comb = ws.cell(r, col_data + 2).value
        quando, precisao = _quando(data, hora)

        combinada = None
        if quando is not None and precisao == "exata" and isinstance(comb, dt.datetime) \
                and comb.replace(second=0, microsecond=0) != quando:
            combinada = comb

        codigo = d.strip().upper() if isinstance(d, str) and RE_CODIGO.match(d.strip().upper()) else None
        principal, extra, tem_extra = _porto(c)
        tonelagem = float(a) if isinstance(a, (int, float)) and isinstance(b, str) \
            and b.strip().lower() == "mt" else None
        rob = float(a) if isinstance(a, (int, float)) and isinstance(b, str) \
            and b.strip().upper() == "ROB" else None
        notas = [n for n in (_nota(ws.cell(r, cc).value)
                             for cc in range(col_data + 3, ws.max_column + 1)) if n]

        if not any((codigo, _acao(e), quando, principal, tem_extra, tonelagem, rob, notas)):
            continue
        linhas.append(_Linha(r, codigo, _acao(e), quando, precisao, notas, principal,
                             extra, tem_extra, tonelagem, rob, combinada))
    return linhas


# ---------------------------------------------------------------------------
# Blocos -> viagens
# ---------------------------------------------------------------------------

@dataclass
class _Bloco:
    linhas: list
    codigo: str | None = None
    porto: str | None = None
    tipo: str | None = None       # alumar_out | faz | juruti | alumar_in | rio


def _blocos(linhas: list[_Linha]) -> list[_Bloco]:
    blocos, atual = [], _Bloco([])
    for linha in linhas:
        if linha.acao == "cabecalho":
            if atual.linhas:
                blocos.append(atual)
            atual = _Bloco([linha])
        else:
            atual.linhas.append(linha)
    if atual.linhas:
        blocos.append(atual)
    for bloco in blocos:
        bloco.codigo = next((l.codigo for l in bloco.linhas if l.codigo), None)
        bloco.porto = next((l.porto_principal for l in bloco.linhas if l.porto_principal), None)
    # Bloco sem codigo herda o do anterior; o primeiro (Rio de Janeiro) pega o do seguinte.
    for i, bloco in enumerate(blocos):
        if bloco.codigo is None:
            anteriores = [b.codigo for b in blocos[:i] if b.codigo]
            seguintes = [b.codigo for b in blocos[i + 1:] if b.codigo]
            bloco.codigo = anteriores[-1] if anteriores else (seguintes[0] if seguintes else None)
    return blocos


def _tipo_do_bloco(bloco: _Bloco, esperado: str) -> str:
    acoes = [l.acao for l in bloco.linhas if l.acao in ("arrival", "berth", "sailing") and l.quando]
    if bloco.porto == "RIO":
        return "rio"
    if bloco.porto == "ALUMAR":
        return "alumar_out" if (acoes and acoes[0] == "sailing") else \
               ("alumar_in" if acoes else esperado)
    if bloco.porto == "FAZENDINHA":
        return "faz"
    if bloco.porto == "JURUTI":
        return "juruti"
    return esperado


def _quantidade_bunker(texto: str) -> tuple[float | None, float | None]:
    """('400 MT VLSFO + 40 MT MGO') -> (400.0, 40.0). Datas e horas sao ignoradas.

    O texto e partido em pedacos ("+", "/", " e "); cada pedaco com um numero
    entre 20 e 1500 e uma quantidade, de MGO se o pedaco disser MGO, senao de
    VLSFO. E o suficiente para todas as grafias que a planilha tem.
    """
    limpo = RE_DATA.sub(" ", texto)
    vlsfo = mgo = None
    for pedaco in re.split(r"\+|/| E ", limpo.upper()):
        for m in RE_NUMERO.finditer(pedaco):
            try:
                valor = float(m.group().replace(",", "."))
            except ValueError:
                continue
            if not 20 <= valor <= 1500:   # 40 t de MGO e pouco; 3 m3 de FW nao e bunker
                continue
            if "MGO" in pedaco:
                mgo = valor if mgo is None else mgo
            else:
                vlsfo = valor if vlsfo is None else vlsfo
            break
    return vlsfo, mgo


def _abastecimento_das_notas(notas: list[str]) -> tuple[tuple | None, bool, str | None]:
    """(abastecimento, cancelado, nota sem quantidade)."""
    for i, nota in enumerate(notas):
        if not RE_BUNKER.search(nota):
            continue
        if re.search(r"CANCEL", nota, re.I):
            return None, True, None
        if re.search(r"VERIFICAR|PDA", nota, re.I) and not RE_NUMERO.search(RE_DATA.sub("", nota)):
            continue
        texto = nota
        vlsfo, mgo = _quantidade_bunker(nota)
        if vlsfo is None and mgo is None:
            # a quantidade pode estar na nota seguinte ("ABSTECIMENTO 19/11" / "400 MT")
            for prox in notas[i + 1:i + 3]:
                v2, m2 = _quantidade_bunker(prox)
                if v2 is not None or m2 is not None:
                    vlsfo, mgo, texto = v2, m2, nota + " / " + prox
                    break
        if vlsfo is None and mgo is None:
            return None, False, nota
        return (vlsfo, mgo, texto), False, None
    return None, False, None


def _codigo_seguinte(codigo: str) -> str:
    """APT26017 -> APT26018. A sequencia e por navio e por ano."""
    return "{}{:03d}".format(codigo[:-3], int(codigo[-3:]) + 1)


def _nova_viagem(navio_id: int, codigo: str, aba: str, linha: int) -> Viagem:
    v = Viagem(navio_id, codigo, aba, linha)
    for chave in CHAVES_MODELO:
        v.escalas[chave] = Escala(chave)
    return v


def _marco(linha: _Linha, rotulo: str | None = None, nota: str | None = None) -> Marco:
    return Marco(linha.quando, linha.n, linha.precisao, rotulo or (linha.acao or ""), nota)


def _parada_extra(linhas: list[_Linha], viagem: Viagem, aba: str, problemas: list) -> Escala | None:
    """As linhas depois do Sailing de Alumar: uma parada de bunker (ou Barra Norte)."""
    com_hora = [l for l in linhas if l.quando]
    notas = [n for l in linhas for n in l.notas]
    abastecimento, cancelado, sem_quantidade = _abastecimento_das_notas(notas)
    if cancelado:
        return None
    if not com_hora:
        return None

    porto = next((l.porto_extra for l in linhas if l.tem_porto_extra and l.porto_extra), None)
    fundeadouro_sem_nome = porto is None
    saida_alumar = viagem.escalas["abertura"].marcos.get("sailing")
    presumido = False
    if porto is None:
        horas = ((com_hora[0].quando - saida_alumar.quando).total_seconds() / 3600
                 if saida_alumar else None)
        porto = "ITAQUI" if horas is not None and horas < HORAS_ITAQUI else "ICOARACI"
        presumido = True
        problemas.append(Problema(
            "aviso", aba, viagem.codigo, com_hora[0].n,
            "parada de bunker sem porto na planilha: presumido {} ({} h depois da saida de Alumar)".format(
                porto, "%.0f" % horas if horas is not None else "?")))

    escala = Escala("extra", linha=com_hora[0].n, porto=porto,
                    motivo="espera_mare" if porto == "BARRA_NORTE" else "bunker",
                    porto_presumido=presumido)
    chegadas = [l for l in com_hora if l.acao in ("arrival", "fundeado")]
    primeira = min(chegadas, key=lambda l: l.quando) if chegadas else com_hora[0]
    escala.marcos["arrival"] = _marco(
        primeira, nota=None if primeira.acao in ("arrival", "fundeado") else
        "Chegada nao registrada na planilha; usada a linha '{}' ({})".format(
            primeira.acao or "sem rotulo", primeira.quando.strftime("%d/%m %H:%M")))

    saidas = [l for l in com_hora if l.acao == "sailing"]
    ultima = com_hora[-1]
    if saidas:
        escala.marcos["sailing"] = _marco(saidas[-1])
    elif ultima is not primeira and ultima.acao in (None, "termino"):
        escala.marcos["sailing"] = _marco(
            ultima, nota="Saida nao registrada na planilha; usada a linha '{}' ({})".format(
                ultima.acao or "sem rotulo", ultima.quando.strftime("%d/%m %H:%M")))

    inicio = next((l for l in com_hora if l.acao == "inicio"), None)
    termino = next((l for l in com_hora if l.acao == "termino"), None)
    if inicio or termino:
        escala.notas.append("Abastecimento: início {} · término {}".format(
            inicio.quando.strftime("%d/%m %H:%M") if inicio else "?",
            termino.quando.strftime("%d/%m %H:%M") if termino else "?"))
    escala.notas.extend(notas)
    escala.abastecimento = abastecimento
    if sem_quantidade:
        problemas.append(Problema("aviso", aba, viagem.codigo, escala.linha,
                                  "nota de abastecimento sem quantidade legível: {!r}".format(sem_quantidade)))
    for l in com_hora:
        if l.acao not in ("arrival", "fundeado", "inicio", "termino", "sailing", "berth", None):
            problemas.append(Problema("nota", aba, viagem.codigo, l.n, "linha ignorada na parada adicional"))
    del fundeadouro_sem_nome
    return escala


def _preencher(bloco: _Bloco, viagem: Viagem, aba: str, problemas: list) -> None:
    tipo = bloco.tipo
    linhas = bloco.linhas

    if tipo == "rio":
        saida = next((l for l in linhas if l.acao == "sailing" and l.quando), None)
        ab = viagem.escalas["abertura"]
        ab.cancelar = "Saída do Rio de Janeiro em {} — viagem de entrega, sem saída de Alumar".format(
            saida.quando.strftime("%d/%m/%Y %H:%M") if saida else "data não registrada")
        ab.notas.extend(n for l in linhas for n in l.notas)
        return

    if tipo == "alumar_out":
        ab = viagem.escalas["abertura"]
        i_sail = next((i for i, l in enumerate(linhas) if l.acao == "sailing" and l.quando), None)
        if i_sail is None:
            # A saida pode vir do bloco de descarga da viagem anterior (encadeamento).
            ab.notas.extend(n for l in linhas for n in l.notas)
            return
        ab.marcos["sailing"] = _marco(linhas[i_sail])
        depois = linhas[i_sail + 1:]
        antes = linhas[:i_sail + 1]
        ab.notas.extend(n for l in antes for n in l.notas)
        so_saidas = bool(depois) and all(l.acao == "sailing" for l in depois if l.quando)
        extra = _parada_extra(depois, viagem, aba, problemas) if depois and not so_saidas else None
        if extra:
            viagem.extras.append(extra)
            if extra.abastecimento is None:
                extra.abastecimento, _c, _s = _abastecimento_das_notas(ab.notas)
        else:
            notas = [n for l in depois for n in l.notas]
            abastecimento, cancelado, sem_qtd = _abastecimento_das_notas(ab.notas + notas)
            ab.notas.extend(notas)
            for l in depois:
                if l.quando and l.acao == "sailing":
                    ab.notas.append("Segunda saída registrada às {}".format(
                        l.quando.strftime("%d/%m/%Y %H:%M")))
            if cancelado:
                ab.notas.append("Abastecimento cancelado (planilha)")
            elif abastecimento:
                ab.abastecimento = abastecimento
            elif sem_qtd:
                problemas.append(Problema("aviso", aba, viagem.codigo, linhas[0].n,
                                          "nota de abastecimento sem quantidade legível: {!r}".format(sem_qtd)))
            for l in depois:
                if l.quando and l.acao != "sailing":
                    problemas.append(Problema("nota", aba, viagem.codigo, l.n,
                                              "linha com data depois da saída de Alumar, ignorada (abastecimento cancelado ou sem parada)"))
        return

    if tipo == "extra":
        extra = _parada_extra(linhas, viagem, aba, problemas)
        if extra:
            if extra.abastecimento is None:
                extra.abastecimento, _c, _s = _abastecimento_das_notas(viagem.escalas["abertura"].notas)
            viagem.extras.append(extra)
        return

    chave = {"faz": "faz_subida", "juruti": "juruti", "alumar_in": "encerramento"}[tipo]
    escala = viagem.escalas[chave]
    escala.linha = linhas[0].n
    aceitos = {"faz": ("arrival", "sailing"),
               "juruti": ("arrival", "berth", "unberth", "sailing"),
               "alumar_in": ("arrival", "berth", "unberth", "sailing")}[tipo]
    i_sail = None
    for i, l in enumerate(linhas):
        if l.tonelagem is not None:
            escala.tonelagem = l.tonelagem
        if l.rob is not None:
            viagem.rob_planilha = l.rob
        if l.combinada is not None:
            problemas.append(Problema(
                "aviso", aba, viagem.codigo, l.n,
                "{} {}: a coluna data+hora diz {}, a coluna combinada diz {} — usada a data+hora".format(
                    escala.rotulo, l.acao or "", l.quando.strftime("%d/%m/%Y %H:%M"),
                    l.combinada.strftime("%d/%m/%Y %H:%M"))))
        if l.quando is None:
            continue
        if i_sail is not None:
            continue                      # depois do Sailing: tratado abaixo
        if l.acao in aceitos:
            if l.acao in escala.marcos:
                problemas.append(Problema("erro", aba, viagem.codigo, l.n,
                                          "{} tem dois {}".format(escala.rotulo, l.acao)))
                continue
            escala.marcos[l.acao] = _marco(l)
            if l.acao == "sailing":
                i_sail = i
            if l.precisao == "apenas_data":
                problemas.append(Problema("aviso", aba, viagem.codigo, l.n,
                                          "{} {}: só a data, sem hora (gravado como 'apenas data')".format(
                                              escala.rotulo, l.acao)))
        else:
            problemas.append(Problema("nota", aba, viagem.codigo, l.n,
                                      "{}: linha '{}' com data fora do padrão, ignorada".format(
                                          escala.rotulo, l.acao or "sem rótulo")))

    notas = [n for l in linhas for n in l.notas]
    escala.notas.extend(notas)

    if tipo == "alumar_in" and "sailing" in escala.marcos:
        viagem.saida_registrada = escala.marcos.pop("sailing")

    # Depois do Sailing de Juruti: a troca de tripulacao na descida (Fazendinha).
    if tipo == "juruti" and i_sail is not None:
        depois = [l for l in linhas[i_sail + 1:] if l.quando]
        faz = viagem.escalas["faz_descida"]
        com_hora = [l for l in depois if l.precisao == "exata"]
        so_data = [l for l in depois if l.precisao != "exata"]
        if len(com_hora) >= 2:
            faz.marcos["arrival"] = _marco(com_hora[0], "sem rótulo",
                                           "Linha sem rótulo na planilha, lida como chegada (troca de tripulação)")
            faz.marcos["sailing"] = _marco(com_hora[-1], "sem rótulo",
                                           "Linha sem rótulo na planilha, lida como saída (troca de tripulação)")
            faz.linha = com_hora[0].n
            problemas.append(Problema("aviso", aba, viagem.codigo, com_hora[0].n,
                                      "par de linhas sem rótulo depois de Juruti lido como Fazendinha na descida (troca de tripulação): {} → {}".format(
                                          com_hora[0].quando.strftime("%d/%m %H:%M"),
                                          com_hora[-1].quando.strftime("%d/%m %H:%M"))))
        elif com_hora:
            problemas.append(Problema("nota", aba, viagem.codigo, com_hora[0].n,
                                      "linha solta com hora depois de Juruti, ignorada"))
        for l in so_data:
            faz.notas.append("Registro na planilha em {} sem hora: {}".format(
                l.quando.strftime("%d/%m/%Y"), " / ".join(l.notas) or "sem texto"))

    # Abastecimento anotado num bloco de porto (esperando berco em Alumar, por exemplo).
    abastecimento, cancelado, sem_qtd = _abastecimento_das_notas(notas)
    if abastecimento:
        escala.abastecimento = abastecimento
    elif sem_qtd:
        problemas.append(Problema("aviso", aba, viagem.codigo, escala.linha,
                                  "nota de abastecimento sem quantidade legível: {!r}".format(sem_qtd)))


# ---------------------------------------------------------------------------
# As decisoes, aplicadas
# ---------------------------------------------------------------------------

NOTA_UNBERTH = ("Unberth não registrado no histórico; copiado do Sailing "
                "(decisão de {}, opção B)".format(DATA_DECISOES))
NOTA_SEM_REGISTRO = "Sem registro no histórico (planilha das supervisoras)"


def _aplicar_decisoes(viagens: list[Viagem], aba: str, problemas: list) -> None:
    for i, v in enumerate(viagens):
        # 1. Unberth = Sailing
        for chave in ("juruti", "encerramento"):
            e = v.escalas[chave]
            saida = e.marcos.get("sailing") if chave == "juruti" else v.saida_registrada
            if "unberth" in e.marcos:
                problemas.append(Problema("nota", aba, v.codigo, e.marcos["unberth"].linha,
                                          "{} tem Unberth registrado de verdade — mantido".format(e.rotulo)))
            elif saida is not None:
                e.marcos["unberth"] = Marco(saida.quando, saida.linha, saida.precisao,
                                            "sailing", NOTA_UNBERTH)

        # Encadeamento: a saida de Alumar desta viagem abre a seguinte.
        if i + 1 < len(viagens):
            prox = viagens[i + 1].escalas["abertura"]
            if v.saida_registrada is None:
                continue
            if "sailing" not in prox.marcos:
                if not prox.cancelar:
                    prox.marcos["sailing"] = Marco(
                        v.saida_registrada.quando, v.saida_registrada.linha,
                        v.saida_registrada.precisao, "sailing",
                        "Saída copiada do bloco de descarga da viagem anterior ({})".format(v.codigo))
            elif prox.marcos["sailing"].quando != v.saida_registrada.quando:
                problemas.append(Problema(
                    "aviso", aba, viagens[i + 1].codigo, prox.marcos["sailing"].linha,
                    "saída de Alumar registrada como {} no fim da {} e {} no início da {} — usada a desta viagem".format(
                        v.saida_registrada.quando.strftime("%d/%m %H:%M"), v.codigo,
                        prox.marcos["sailing"].quando.strftime("%d/%m %H:%M"), viagens[i + 1].codigo)))
                prox.notas.append("Planilha: saída registrada como {} na viagem anterior".format(
                    v.saida_registrada.quando.strftime("%d/%m/%Y %H:%M")))

    # 2. Escalas sem registro nascem canceladas
    for v in viagens:
        passou_de_alumar = bool(v.escalas["encerramento"].marcos)
        for chave in CHAVES_MODELO:
            e = v.escalas[chave]
            if e.cancelar or e.marcos:
                continue
            if not v.aberta or (chave in ("faz_descida", "barra_norte") and passou_de_alumar) \
                    or (chave == "faz_subida" and v.escalas["juruti"].marcos):
                e.cancelar = NOTA_SEM_REGISTRO


def _validar(viagens: list[Viagem], aba: str, navio_id: int, problemas: list) -> None:
    vistos = set()
    saldo = 0.0
    for v in viagens:
        if not v.codigo.startswith(PREFIXOS[navio_id]):
            problemas.append(Problema("erro", aba, v.codigo, v.linha,
                                      "código não é do navio desta aba (prefixo {})".format(PREFIXOS[navio_id])))
        if v.codigo in vistos:
            problemas.append(Problema("erro", aba, v.codigo, v.linha, "código repetido"))
        vistos.add(v.codigo)

        # ordem cronologica de tudo que vai ser gravado
        anterior = None
        for escala, tipo, marco in v.todos_os_marcos():
            if anterior and marco.quando < anterior[2].quando:
                problemas.append(Problema(
                    "erro", aba, v.codigo, marco.linha,
                    "{} {} ({}) é anterior a {} {} ({})".format(
                        escala.rotulo, tipo, marco.quando.strftime("%d/%m %H:%M"),
                        anterior[0].rotulo, anterior[1], anterior[2].quando.strftime("%d/%m %H:%M"))))
            anterior = (escala, tipo, marco)

        ab = v.escalas["abertura"]
        if "sailing" not in ab.marcos and not ab.cancelar and any(True for _ in v.todos_os_marcos()):
            problemas.append(Problema("erro", aba, v.codigo, v.linha,
                                      "viagem sem saída de Alumar e sem saída registrada na viagem anterior"))
        for e in v.extras:
            if e.motivo == "bunker" and e.abastecimento is None:
                problemas.append(Problema("aviso", aba, v.codigo, e.linha,
                                          "{} sem nota de quantidade abastecida".format(e.rotulo)))
        enc = v.escalas["encerramento"]
        if not v.aberta and "unberth" not in enc.marcos:
            problemas.append(Problema("erro", aba, v.codigo, v.linha,
                                      "viagem encerrada sem saída de Alumar: não há como fechá-la"))
        if not v.aberta and not v.escalas["juruti"].marcos:
            problemas.append(Problema("aviso", aba, v.codigo, v.linha, "viagem sem nenhum marco em Juruti"))
        if not v.aberta and v.escalas["faz_subida"].cancelar == NOTA_SEM_REGISTRO:
            problemas.append(Problema("aviso", aba, v.codigo, v.linha,
                                      "Fazendinha na subida sem registro nenhum — cancelada"))
        for e in v.extras:
            if "sailing" not in e.marcos:
                problemas.append(Problema("aviso", aba, v.codigo, e.linha,
                                          "{} sem saída registrada".format(e.rotulo)))
        # Marco que a planilha deixou vazio numa viagem encerrada: fica em falta
        # na tela de Encerradas, e e verdade — mas o relatorio avisa.
        exigidos = {"abertura": ("sailing",), "faz_subida": ("arrival", "sailing"),
                    "juruti": ("arrival", "berth", "unberth", "sailing"),
                    "faz_descida": ("arrival", "sailing"), "barra_norte": ("arrival", "sailing"),
                    "encerramento": ("arrival", "berth", "unberth")}
        if not v.aberta:
            for chave, tipos in exigidos.items():
                e = v.escalas[chave]
                faltam = [t for t in tipos if t not in e.marcos] if not e.cancelar else []
                if faltam:
                    problemas.append(Problema("aviso", aba, v.codigo, e.linha or v.linha,
                                              "{} sem {} — fica como marco em falta".format(
                                                  e.rotulo, ", ".join(faltam))))

        # O ROB da planilha e o SALDO a bordo, acumulado de viagem em viagem —
        # a mesma conta da view carga_bordo. Divergencia aqui e conferencia, nao erro.
        saldo += (v.escalas["juruti"].tonelagem or 0) - (enc.tonelagem or 0)
        if v.rob_planilha is not None and abs(saldo - v.rob_planilha) > 0.5:
            problemas.append(Problema("aviso", aba, v.codigo, v.linha,
                                      "ROB da planilha ({:g}) difere do saldo acumulado ({:g})".format(
                                          v.rob_planilha, saldo)))
            saldo = v.rob_planilha


def _marcos_manuais(viagens: list[Viagem], aba: str, problemas: list) -> None:
    for aba_alvo, codigo, chave, tipo, iso, motivo in MARCOS_MANUAIS:
        if aba_alvo != aba:
            continue
        v = next((x for x in viagens if x.codigo == codigo), None)
        escala = None
        if v is not None:
            if chave.startswith("extra:"):
                escala = next((e for e in v.extras if e.porto == chave[6:]), None)
            else:
                escala = v.escalas.get(chave)
        if escala is None:
            problemas.append(Problema("aviso", aba, codigo, None,
                                      "marco manual NÃO aplicado, escala não encontrada: {}".format(motivo)))
            continue
        escala.marcos[tipo] = Marco(dt.datetime.fromisoformat(iso), 0, "exata", "manual", motivo)
        problemas.append(Problema("nota", aba, codigo, None, "marco manual aplicado: {}".format(motivo)))


def ler(caminho) -> tuple[dict[int, list[Viagem]], list[Problema]]:
    """navio_id -> viagens na ordem, e a lista de problemas encontrados."""
    import openpyxl

    wb = openpyxl.load_workbook(caminho, data_only=True)
    resultado, problemas = {}, []
    for ws in wb.worksheets:
        aba = _norm(ws.title)
        navio_id = NAVIOS.get(aba)
        if navio_id is None:
            problemas.append(Problema("aviso", aba, None, None, "aba desconhecida, ignorada"))
            continue
        _aplicar_ajustes(ws, aba, problemas)
        linhas = _ler_linhas(ws, aba, problemas)
        blocos = _blocos(linhas)

        viagens: list[Viagem] = []
        por_codigo: dict[str, Viagem] = {}
        contagem: dict[str, int] = {}
        ordem_esperada = ("alumar_out", "faz", "juruti", "alumar_in")
        for bloco in blocos:
            if bloco.codigo is None:
                for l in bloco.linhas:
                    if l.quando or l.notas:
                        problemas.append(Problema("nota", aba, None, l.n, "linha fora de qualquer viagem, ignorada"))
                continue
            v = por_codigo.get(bloco.codigo)
            if v is None:
                v = _nova_viagem(navio_id, bloco.codigo, aba, bloco.linhas[0].n)
                por_codigo[bloco.codigo] = v
                viagens.append(v)
                contagem[bloco.codigo] = 0
            n = contagem[bloco.codigo]
            esperado = ordem_esperada[min(n, 3)]
            bloco.tipo = _tipo_do_bloco(bloco, esperado)
            e_parada = bloco.porto is None and n == 1 and any(
                l.tem_porto_extra or l.acao in ("fundeado", "inicio", "termino") for l in bloco.linhas)
            if e_parada:
                bloco.tipo = "extra"
            elif bloco.tipo == "rio":
                contagem[bloco.codigo] += 1
            else:
                if bloco.porto is None:
                    problemas.append(Problema("nota", aba, v.codigo, bloco.linhas[0].n,
                                              "bloco sem porto na coluna C; lido como {} pela posição".format(bloco.tipo)))
                elif bloco.tipo != esperado and n < 4:
                    problemas.append(Problema("aviso", aba, v.codigo, bloco.linhas[0].n,
                                              "bloco de {} onde se esperava {}".format(bloco.tipo, esperado)))
                contagem[bloco.codigo] += 1
            _preencher(bloco, v, aba, problemas)

        vazias = [v for v in viagens if not any(e.marcos for e in v.escalas.values())
                  and not v.extras]
        for v in vazias:
            problemas.append(Problema("nota", aba, v.codigo, v.linha, "viagem sem nenhum marco, não entra"))
            viagens.remove(v)
        if viagens and viagens[-1].saida_registrada is not None:
            # A ultima viagem da aba ja saiu de Alumar: esta encerrada, e o navio
            # esta na seguinte — que o sistema abriria sozinho no Unberth. Nasce
            # aqui, vazia, com a saida herdada (o encadeamento faz isso abaixo).
            ultima = viagens[-1]
            codigo = vazias[0].codigo if vazias else _codigo_seguinte(ultima.codigo)
            problemas.append(Problema("nota", aba, codigo, ultima.linha,
                                      "viagem aberta pelo importador: a {} já saiu de Alumar e a planilha "
                                      "não tem a seguinte".format(ultima.codigo)))
            viagens.append(_nova_viagem(navio_id, codigo, aba, ultima.linha))
        if viagens:
            viagens[-1].aberta = True
        _marcos_manuais(viagens, aba, problemas)
        _aplicar_decisoes(viagens, aba, problemas)
        _validar(viagens, aba, navio_id, problemas)
        resultado[navio_id] = viagens
    return resultado, problemas


def resumo(viagens_por_navio: dict[int, list[Viagem]]) -> list[dict]:
    """Uma linha por navio, para o relatorio."""
    saida = []
    for navio_id, viagens in sorted(viagens_por_navio.items()):
        marcos = sum(1 for v in viagens for _ in v.todos_os_marcos())
        copiados = sum(1 for v in viagens for _e, _t, m in v.todos_os_marcos() if m.nota == NOTA_UNBERTH)
        saida.append({
            "navio_id": navio_id, "prefixo": PREFIXOS[navio_id], "viagens": len(viagens),
            "primeira": viagens[0].codigo if viagens else "", "ultima": viagens[-1].codigo if viagens else "",
            "de": min((m.quando for v in viagens for _e, _t, m in v.todos_os_marcos()), default=None),
            "ate": max((m.quando for v in viagens for _e, _t, m in v.todos_os_marcos()), default=None),
            "marcos": marcos, "unberth_copiados": copiados,
            "extras": sum(len(v.extras) for v in viagens),
            "abastecimentos": sum(1 for v in viagens for e in list(v.escalas.values()) + v.extras
                                  if e.abastecimento),
            "cargas": sum(1 for v in viagens for e in v.escalas.values() if e.tonelagem is not None),
            "canceladas": sum(1 for v in viagens for e in v.escalas.values() if e.cancelar),
        })
    return saida
