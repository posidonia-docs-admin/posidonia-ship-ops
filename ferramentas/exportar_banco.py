# -*- coding: utf-8 -*-
"""Exporta o banco do Corsair para um Excel: uma aba por assunto.

    python -X utf8 -m ferramentas.exportar_banco [--saida caminho.xlsx] [--env .env]

Abas: Viagens (uma linha por viagem, com carga, duração e consumo), Pernadas
(horas de cada pernada de cada viagem), Escalas (os quatro marcos em colunas),
Marcos (todo evento vigente), Abastecimentos, Carga. O banco é o que
app.db.conectar() decide: local por padrão, Turso com --env.
"""
from __future__ import annotations

import argparse
import os
import sys
from contextlib import closing
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

if "--env" in sys.argv:
    from ferramentas.importar_historico import carregar_env
    carregar_env(sys.argv[sys.argv.index("--env") + 1])
os.environ.setdefault("SHIPOPS_SECRET_KEY", "sem-uso-na-exportacao")

from app import analises, db, pernadas  # noqa: E402

NAVY = "0A2540"


def _data(texto):
    """'2026-03-01T18:40' -> datetime, para o Excel tratar como data."""
    if not texto:
        return None
    try:
        return datetime.fromisoformat(texto[:16])
    except ValueError:
        return texto


def _aba(wb, nome, cabecalho, linhas, larguras=None, formatos=None):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(nome)
    ws.append(cabecalho)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center")
    for linha in linhas:
        ws.append(list(linha))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, _ in enumerate(cabecalho, start=1):
        largura = (larguras or {}).get(i - 1, 14)
        ws.column_dimensions[get_column_letter(i)].width = largura
    for coluna, formato in (formatos or {}).items():
        for celula in ws.iter_cols(min_col=coluna + 1, max_col=coluna + 1, min_row=2):
            for c in celula:
                c.number_format = formato
    return ws


def exportar(conn, saida: Path) -> dict:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    contagens = {}

    # ---- Viagens: a base de Análises, uma linha por viagem
    base = analises.carregar(conn, "viagens")
    inicio_fim = {r["id"]: (r["inicio"], r["termino"], r["status"]) for r in conn.execute(
        "SELECT vg.id, vg.status, ab.hora_local AS inicio, en.hora_local AS termino "
        "  FROM viagem vg LEFT JOIN evento ab ON ab.id = vg.evento_abertura_id "
        "  LEFT JOIN evento en ON en.id = vg.evento_encerramento_id")}
    numeros = {r[0]: r[1] for r in conn.execute("SELECT id, numero FROM viagem")}
    linhas = []
    for v in base:
        d, m = v["dims"], v["medidas"]
        ini, fim, status = inicio_fim[v["viagem_id"]]
        linhas.append([
            d["navio"][1], numeros[v["viagem_id"]],
            status, _data(ini), _data(fim),
            None if d["ano"] == analises.EM_CURSO else int(d["ano"][1]),
            None if d["mes"] == analises.EM_CURSO else d["mes"][1],
            d["bunker"][1], d["adicional"][1],
            m["carregado"], m["descarregado"], m["duracao"], m["espera_berco"], m["atracado"],
            m["vlsfo_consumido"], m["mgo_consumido"], m["vlsfo_abastecido"], m["mgo_abastecido"],
        ])
    _aba(wb, "Viagens",
         ["Navio", "Viagem", "Situação", "Saída de Alumar", "Encerramento (Unberth Alumar)",
          "Ano", "Mês", "Teve bunker", "Parada adicional", "Carregado (MT)", "Descarregado (MT)",
          "Duração (h)", "Espera de berço (h)", "Tempo atracado (h)", "VLSFO consumido (t)",
          "MGO consumido (t)", "VLSFO abastecido (t)", "MGO abastecido (t)"],
         linhas, {0: 20, 3: 18, 4: 26}, {3: "dd/mm/yyyy hh:mm", 4: "dd/mm/yyyy hh:mm",
                                        9: "#,##0", 10: "#,##0", 11: "0.0", 12: "0.0", 13: "0.0",
                                        14: "#,##0.000", 15: "#,##0.000", 16: "#,##0.000", 17: "#,##0.000"})
    contagens["Viagens"] = len(linhas)

    # ---- Pernadas: horas de cada pernada de cada viagem
    per = analises.carregar(conn, "pernadas")
    linhas = [[l["dims"]["navio"][1], l["dims"]["viagem"][1], l["dims"]["situacao"][1],
               None if l["dims"]["mes"] == analises.EM_CURSO else l["dims"]["mes"][1],
               l["dims"]["grupo"][1], l["dims"]["pernada"][1], l["medidas"]["horas"],
               l["medidas"]["desvio"]] for l in per]
    _aba(wb, "Pernadas", ["Navio", "Viagem", "Situação", "Mês", "Grupo", "Pernada", "Horas",
                          "Desvio da premissa (h)"], linhas, {0: 20, 5: 36}, {6: "0.0", 7: "0.0"})
    contagens["Pernadas"] = len(linhas)

    # ---- Escalas: os quatro marcos em colunas
    horas_locais = {}
    for r in conn.execute("SELECT escala_id, tipo, hora_local FROM evento_vigente"):
        horas_locais[(r[0], r[1])] = r[2]
    linhas = []
    for r in conn.execute(
            "SELECT e.id AS escala_id, n.nome_oficial, vg.numero, e.ordem, p.nome AS porto, e.tipo_escala, e.sentido, "
            "       e.motivo, e.condicao, e.origem, e.status, "
            "       m.arrival_utc, m.berth_utc, m.unberth_utc, m.sailing_utc, "
            "       ec.horas_espera_berco, ec.horas_atracado, ec.horas_total_escala, e.observacao "
            "  FROM escala e JOIN viagem vg ON vg.id = e.viagem_id JOIN navio n ON n.id = vg.navio_id "
            "  JOIN porto p ON p.codigo = e.codigo_porto "
            "  JOIN escala_marcos m ON m.escala_id = e.id "
            "  JOIN escala_completa ec ON ec.escala_id = e.id "
            " ORDER BY vg.navio_id, vg.id, e.ordem"):
        local = {tipo: _data(horas_locais.get((r["escala_id"], tipo)))
                 for tipo in ("arrival", "berth", "unberth", "sailing")}
        linhas.append([r["nome_oficial"].title(), r["numero"], r["ordem"], r["porto"], r["tipo_escala"],
                       r["sentido"], r["motivo"], r["condicao"], r["origem"], r["status"],
                       local["arrival"], local["berth"], local["unberth"], local["sailing"],
                       r["horas_espera_berco"], r["horas_atracado"], r["horas_total_escala"],
                       r["observacao"]])
    dt = "dd/mm/yyyy hh:mm"
    _aba(wb, "Escalas", ["Navio", "Viagem", "Ordem", "Porto", "Tipo", "Sentido", "Motivo", "Condição",
                         "Origem", "Situação", "Arrival", "Berth", "Unberth", "Sailing",
                         "Espera de berço (h)", "Atracado (h)", "Total na escala (h)", "Observação"],
         linhas, {0: 20, 3: 14, 10: 17, 11: 17, 12: 17, 13: 17, 17: 60},
         {10: dt, 11: dt, 12: dt, 13: dt, 14: "0.00", 15: "0.00", 16: "0.00"})
    contagens["Escalas"] = len(linhas)

    # ---- Marcos: todo evento vigente
    linhas = [[r["nome_oficial"].title(), r["numero"], r["porto"], r["ordem"], r["tipo"],
               _data(r["hora_local"]), r["offset_utc"], _data(r["hora_utc"]), r["precisao"],
               r["rob_vlsfo"], r["rob_mgo"], r["nome_responsavel"], r["versao"], r["observacao"]]
              for r in conn.execute(
            "SELECT n.nome_oficial, vg.numero, p.nome AS porto, e.ordem, ev.tipo, ev.hora_local, "
            "       ev.offset_utc, ev.hora_utc, ev.precisao, ev.rob_vlsfo, ev.rob_mgo, "
            "       ev.nome_responsavel, ev.versao, ev.observacao "
            "  FROM evento_vigente ev JOIN escala e ON e.id = ev.escala_id "
            "  JOIN viagem vg ON vg.id = e.viagem_id JOIN navio n ON n.id = vg.navio_id "
            "  JOIN porto p ON p.codigo = e.codigo_porto "
            " ORDER BY vg.navio_id, ev.hora_utc, e.ordem")]
    _aba(wb, "Marcos", ["Navio", "Viagem", "Porto", "Ordem da escala", "Marco", "Hora local", "Fuso",
                        "Hora UTC", "Precisão", "ROB VLSFO (t)", "ROB MGO (t)", "Quem lançou", "Versão",
                        "Observação"], linhas, {0: 20, 5: 17, 7: 17, 11: 34, 13: 70},
         {5: dt, 7: dt, 9: "#,##0.000", 10: "#,##0.000"})
    contagens["Marcos"] = len(linhas)

    # ---- Abastecimentos
    linhas = [[r["nome_oficial"].title(), r["numero"], r["porto"], r["origem"], r["vlsfo"], r["mgo"],
               _data(r["quando"]), r["observacao"]] for r in conn.execute(
            "SELECT n.nome_oficial, vg.numero, p.nome AS porto, e.origem, a.vlsfo, a.mgo, a.observacao, "
            "       (SELECT MAX(hora_local) FROM evento_vigente x WHERE x.escala_id = e.id) AS quando "
            "  FROM abastecimento a JOIN escala e ON e.id = a.escala_id "
            "  JOIN viagem vg ON vg.id = e.viagem_id JOIN navio n ON n.id = vg.navio_id "
            "  JOIN porto p ON p.codigo = e.codigo_porto ORDER BY vg.navio_id, quando")]
    _aba(wb, "Abastecimentos", ["Navio", "Viagem", "Porto", "Parada", "VLSFO (t)", "MGO (t)",
                                "Último marco da parada", "Nota da planilha"], linhas,
         {0: 20, 6: 20, 7: 60}, {4: "#,##0.000", 5: "#,##0.000", 6: dt})
    contagens["Abastecimentos"] = len(linhas)

    # ---- Carga
    linhas = [[r["navio"].title(), r["viagem"], r["porto"], r["condicao"], r["carregado"],
               r["descarregado"], r["carga_bordo"], _data(r["momento"])] for r in conn.execute(
            "SELECT * FROM carga_bordo ORDER BY navio_id, momento, escala_id")]
    _aba(wb, "Carga", ["Navio", "Viagem", "Porto", "Condição", "Carregado (MT)", "Descarregado (MT)",
                       "Saldo a bordo (MT)", "Momento"], linhas, {0: 20, 7: 17},
         {4: "#,##0", 5: "#,##0", 6: "#,##0", 7: dt})
    contagens["Carga"] = len(linhas)

    # ---- Premissas
    orcado = pernadas.premissas(conn)
    nomes = {p[0]: p[2] for p in pernadas.PERNADAS}
    _aba(wb, "Premissas", ["Pernada", "Horas orçadas"],
         [[nomes[k], v] for k, v in orcado.items()], {0: 36}, {1: "0.0"})

    wb.save(saida)
    return contagens


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--saida", default=None)
    parser.add_argument("--env")
    args = parser.parse_args(argv)
    saida = Path(args.saida) if args.saida else (
        Path(__file__).resolve().parent.parent / "Corsair - Banco {}.xlsx".format(db.agora()[:10]))
    with closing(db.conectar()) as conn:
        contagens = exportar(conn, saida)
    print("Gravado em {}".format(saida))
    for aba, n in contagens.items():
        print("  {:<16} {:>6} linhas".format(aba, n))
    return 0


if __name__ == "__main__":
    sys.exit(main())
