# -*- coding: utf-8 -*-
"""O importador do historico: le a planilha das supervisoras e grava no modelo.

A planilha de teste e montada aqui, no layout real (uma aba por navio, blocos
por parada, coluna F = data, G = hora). Cobre o que a planilha de verdade tem:
o Unberth que nao existe, a parada de bunker depois da saida de Alumar, a troca
de tripulacao na descida, as notas, a tonelagem e a viagem em curso.
"""
import datetime as dt

import openpyxl
import pytest

from app import pernadas
from ferramentas import historico_planilha as hp
from ferramentas import importar_historico as imp


def _planilha(caminho, *, com_erro=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AMAZON PIONEER"
    r = [1]

    def linha(a=None, b=None, c=None, d=None, e=None, data=None, hora=None, **notas):
        n = r[0]
        for col, v in zip("ABCDE", (a, b, c, d, e)):
            if v is not None:
                ws[f"{col}{n}"] = v
        if data is not None:
            ws[f"F{n}"] = dt.datetime(*data)
        if hora is not None:
            ws[f"G{n}"] = dt.time(*hora)
        for col, v in notas.items():
            ws[f"{col}{n}"] = v
        r[0] += 1
        return n

    def cabecalho(porto=None, codigo=None):
        return linha(c=porto, d=codigo or "VOY", e="ACTION")

    # ---- APN24001: viagem completa, com bunker na espera de berco em Alumar
    cabecalho("ALUMAR")
    linha(a="ORION")
    linha(d="APN24001", e="SAILING", data=(2024, 11, 23), hora=(15, 0))
    cabecalho("FAZENDINHA")
    linha(d="APN24001", e="ARRIVAL", data=(2024, 11, 25), hora=(15, 0), K="PDA RECEBIDA 27/11/2024")
    linha(a="ORION", d="APN24001", e="BERTH")
    linha(a="PASSAGEM", d="APN24001", e="UNBERTH")
    linha(d="APN24001", e="SAILING", data=(2024, 11, 28), hora=(2, 20))
    cabecalho("JURUTI")
    linha(d="APN24001", e="ARRIVAL", data=(2024, 11, 29), hora=(11, 5))
    linha(a="ORION", d="APN24001", e="BERTH", data=(2024, 11, 29), hora=(18, 33))
    linha(a="LOADING", d="APN24001", e="UNBERTH")
    linha(a=53960, b="mt", d="APN24001", e="SAILING", data=(2024, 11, 30), hora=(17, 15))
    cabecalho("ALUMAR  -MA ", "APN24001")
    linha(d="APN24001", e="ARRIVAL", data=(2024, 12, 4), hora=(9, 30), K="*** TROCA DE TRIPULAÇÃO 05/12/2024 ")
    linha(a="ORION", d="APN24001", e="BERTH", data=(2024, 12, 10), hora=(15, 40))
    linha(a="DISCHARGING", d="APN24001", e="UNBERTH")
    linha(a=53906, b="mt", d="APN24001", e="SAILING", data=(2024, 12, 12), hora=(2, 55))
    linha(K="08/12 - ABASTECIMENTO - VLSFO - 350 mt")
    linha(a=54, b="ROB", e="Duração da Viagem")
    linha()

    # ---- APN24002: parada de bunker em Icoaraci (FUNDEADO/INICIO/TERMINO/SAILING)
    #      e troca de tripulacao na descida (par sem rotulo depois de Juruti)
    cabecalho("ALUMAR  -MA ")
    linha(a="ORION")
    linha(d="APN24002", e="SAILING", data=(2024, 12, 12), hora=(2, 55))
    linha(d="APN24002", e="ARRIVAL", data=(2024, 12, 13), hora=(10, 0))
    linha(c="MOSQUEIRO", d="APN24002", e="FUNDEADO", data=(2024, 12, 13), hora=(11, 0))
    linha(c="ICOARACI", d="APN24002", e="Inicio Operação", data=(2024, 12, 13), hora=(14, 0),
          K="394,650 mt VLSFO + 55,574 MT MGO")
    linha(d="APN24002", e="Termino operação", data=(2024, 12, 13), hora=(20, 0))
    linha(d="APN24002", e="SAILING", data=(2024, 12, 13), hora=(22, 0))
    cabecalho("FAZENDINHA")
    linha(d="APN24002", e="ARRIVAL", data=(2024, 12, 15), hora=(13, 50))
    linha(a="ORION", d="APN24002", e="BERTH")
    linha(a="PASSAGEM", d="APN24002", e="UNBERTH")
    linha(d="APN24002", e="SAILING", data=(2024, 12, 15), hora=(14, 0))
    cabecalho("JURUTI")
    linha(a="ORION", d="APN24002", e="ARRIVAL", data=(2024, 12, 16), hora=(21, 55))
    linha(a="LOADING", d="APN24002", e="BERTH", data=(2024, 12, 16), hora=(23, 55))
    linha(a=58166, b="mt", d="APN24002", e="UNBERTH")
    linha(d="APN24002", e="SAILING", data=(2024, 12, 18), hora=(7, 52))
    linha(data=(2024, 12, 20), hora=(5, 30), K="TROCA DE TRIPULAÇÃO")
    linha(data=(2024, 12, 20), hora=(8, 0))
    cabecalho("ALUMAR  -MA ")
    linha(d="APN24002", e="ARRIVAL", data=(2024, 12, 22), hora=(5, 0))
    linha(a="ORION", d="APN24002", e="BERTH", data=(2024, 12, 28), hora=(17, 27))
    linha(a="DISCHARGING", d="APN24002", e="UNBERTH")
    linha(a=56781, b="mt", d="APN24002", e="SAILING", data=(2024, 12, 29), hora=(4, 47))
    linha(a=1439, b="ROB", e="Duração da Viagem")
    linha()

    # ---- APN24003: em curso (chegou em Juruti). A saida nao foi digitada aqui:
    #      vem do bloco de descarga da APN24002.
    cabecalho("ALUMAR  -MA ")
    linha(a="ORION")
    linha(d="APN24003", e="SAILING")
    cabecalho("FAZENDINHA")
    linha(d="APN24003", e="ARRIVAL", data=(2024, 12, 31), hora=(4, 10) if not com_erro else (4, 10))
    linha(a="ORION", d="APN24003", e="BERTH")
    linha(a="PASSAGEM", d="APN24003", e="UNBERTH")
    linha(d="APN24003", e="SAILING", data=(2024, 12, 30) if com_erro else (2024, 12, 31), hora=(5, 0))
    cabecalho("JURUTI")
    linha(a="ORION", d="APN24003", e="ARRIVAL", data=(2025, 1, 2), hora=(11, 35))
    linha(a="LOADING", d="APN24003", e="BERTH")
    linha(d="APN24003", e="UNBERTH")
    linha(d="APN24003", e="SAILING")
    cabecalho("ALUMAR  -MA ")
    linha(d="APN24003", e="ARRIVAL")
    linha(d="APN24003", e="BERTH")
    # ---- APN24004: so linhas planejadas, sem marco nenhum — nao entra
    cabecalho("ALUMAR  -MA ")
    linha(d="APN24004", e="SAILING")

    wb.save(caminho)
    return caminho


@pytest.fixture()
def planilha(tmp_path):
    return _planilha(tmp_path / "historico.xlsx")


def test_le_as_viagens_no_modelo_do_corsair(planilha):
    por_navio, problemas = hp.ler(planilha)
    assert list(por_navio) == [2]                              # Pioneer
    viagens = por_navio[2]
    assert [v.codigo for v in viagens] == ["APN24001", "APN24002", "APN24003"]
    assert not [p for p in problemas if p.nivel == "erro"], problemas

    v1 = viagens[0]
    assert v1.escalas["abertura"].marcos["sailing"].iso == "2024-11-23T15:00"
    juruti = v1.escalas["juruti"]
    assert juruti.marcos["unberth"].iso == "2024-11-30T17:15"          # copiado do Sailing
    assert juruti.marcos["unberth"].nota == hp.NOTA_UNBERTH
    assert juruti.tonelagem == 53960
    alumar = v1.escalas["encerramento"]
    assert alumar.marcos["unberth"].iso == "2024-12-12T02:55"
    assert "sailing" not in alumar.marcos                              # a saida abre a proxima
    assert alumar.tonelagem == 53906
    assert alumar.abastecimento[:2] == (350.0, None)                   # bunker na espera de berco
    assert any("TROCA DE TRIPULA" in n for n in alumar.notas)
    assert v1.escalas["faz_descida"].cancelar == hp.NOTA_SEM_REGISTRO
    assert v1.escalas["barra_norte"].cancelar == hp.NOTA_SEM_REGISTRO
    assert v1.rob_planilha == 54

    v2 = viagens[1]
    assert v2.escalas["abertura"].marcos["sailing"].iso == "2024-12-12T02:55"
    (extra,) = v2.extras
    assert extra.porto == "ICOARACI" and extra.motivo == "bunker"
    assert extra.marcos["arrival"].iso == "2024-12-13T10:00"           # o ARRIVAL, nao o FUNDEADO
    assert extra.marcos["sailing"].iso == "2024-12-13T22:00"
    assert extra.abastecimento[:2] == (394.65, 55.574)
    assert any("início 13/12 14:00" in n for n in extra.notas)
    faz = v2.escalas["faz_descida"]
    assert faz.cancelar is None
    assert faz.marcos["arrival"].iso == "2024-12-20T05:30"
    assert faz.marcos["sailing"].iso == "2024-12-20T08:00"
    assert v2.escalas["barra_norte"].cancelar == hp.NOTA_SEM_REGISTRO

    v3 = viagens[2]
    assert v3.aberta
    assert v3.escalas["abertura"].marcos["sailing"].iso == "2024-12-29T04:47"   # herdada
    assert v3.escalas["faz_descida"].cancelar is None                  # ainda vai acontecer
    assert "berth" not in v3.escalas["juruti"].marcos
    assert any("APN24004" in (p.codigo or "") and "não entra" in p.texto for p in problemas)


def test_navio_que_ja_saiu_de_alumar_ganha_a_viagem_seguinte_aberta(tmp_path):
    """Se a ultima viagem da aba tem o Sailing de Alumar, ela esta encerrada e o
    navio esta na seguinte — que nasce aberta, com a saida herdada, como o sistema
    faria no Unberth."""
    caminho = _planilha(tmp_path / "saiu.xlsx")
    wb = openpyxl.load_workbook(caminho)
    ws = wb.active
    # apaga tudo da APN24003 em diante: a aba termina na APN24002, que ja saiu
    primeira = next(r for r in range(1, ws.max_row + 1) if ws[f"D{r}"].value == "APN24003") - 2
    ws.delete_rows(primeira, ws.max_row)
    wb.save(caminho)

    por_navio, problemas = hp.ler(caminho)
    viagens = por_navio[2]
    assert [v.codigo for v in viagens] == ["APN24001", "APN24002", "APN24003"]
    assert not viagens[1].aberta and viagens[2].aberta
    assert viagens[2].escalas["abertura"].marcos["sailing"].iso == "2024-12-29T04:47"
    assert not any(e.cancelar for e in viagens[2].escalas.values())
    assert any("aberta pelo importador" in p.texto for p in problemas)


def test_marco_fora_de_ordem_e_erro(tmp_path):
    por_navio, problemas = hp.ler(_planilha(tmp_path / "erro.xlsx", com_erro=True))
    erros = [p for p in problemas if p.nivel == "erro"]
    assert len(erros) == 1 and "anterior" in erros[0].texto and erros[0].codigo == "APN24003"


def test_quantidade_de_bunker_em_todas_as_grafias():
    q = hp._quantidade_bunker
    assert q("08/12 - ABASTECIMENTO - VLSFO - 350 mt") == (350.0, None)
    assert q("ABASTECIMENTO 400 MT 04/01/2025") == (400.0, None)
    assert q("400 Vlsfo + 80 MGO") == (400.0, 80.0)
    assert q("Abastecimento VLSFO 400 mt + MGO 40mt") == (400.0, 40.0)
    assert q("600 MT VLSFO+ 40 MT MGO") == (600.0, 40.0)
    assert q("ABASTECIMENTO 500.087 MT ") == (500.087, None)
    assert q("Vlsfo 600") == (600.0, None)
    assert q("11/02 08:00 a 12/02 15:00 manutenção") == (None, None)
    assert hp._abastecimento_das_notas(["CANCELADO-ABASTECIMENTO  400 MT "]) == (None, True, None)
    assert hp._abastecimento_das_notas(["ABSTECIMENTO 19/11/2025", "400 MT"])[0][:2] == (400.0, None)


def test_importa_no_banco_e_as_telas_leem(conn, planilha):
    por_navio, _ = hp.ler(planilha)
    resultado = imp.importar(conn, por_navio, apagar=False)
    assert resultado["viagens"] == 3

    linhas = conn.execute(
        "SELECT numero, status FROM viagem WHERE navio_id = 2 ORDER BY id").fetchall()
    assert [tuple(l) for l in linhas] == [("APN24001", "encerrada"), ("APN24002", "encerrada"),
                                          ("APN24003", "aberta")]
    completa = {r["numero"]: r for r in conn.execute("SELECT * FROM viagem_completa")}
    assert completa["APN24001"]["horas_viagem"] == pytest.approx(18 * 24 + 11.92, abs=0.02)
    assert completa["APN24001"]["marcos_faltantes"] == 0          # canceladas nao cobram
    assert completa["APN24002"]["escalas_extras"] == 1

    # todo marco importado ja esta conferido, por "importacao"
    assert conn.execute("SELECT COUNT(*) FROM escalas_a_conferir").fetchone()[0] == 0
    assert conn.execute("SELECT COUNT(*) FROM evento").fetchone()[0] == \
        conn.execute("SELECT COUNT(*) FROM conferencia WHERE conferido_por = 'importacao'").fetchone()[0]
    assert conn.execute("SELECT ativo FROM conta WHERE login = 'importacao'").fetchone()[0] == 0

    # as pernadas saem da viagem importada como sairiam de uma lancada pelo comandante
    v1 = conn.execute("SELECT id FROM viagem WHERE numero = 'APN24001'").fetchone()[0]
    calc = pernadas.calcular(conn, v1)
    assert calc["sub_nav_alumar_faz"] == 48 * 60
    assert calc["jur_carregamento"] == (22 * 60 + 42)              # Berth -> Unberth(=Sailing)
    assert calc["alu_espera_berco"] == 6 * 24 * 60 + 6 * 60 + 10
    assert calc["des_espera_faz"] is None                          # cancelada: sem dado, nao zero
    v2 = conn.execute("SELECT id FROM viagem WHERE numero = 'APN24002'").fetchone()[0]
    calc2 = pernadas.calcular(conn, v2)
    assert calc2["sub_bunker_saida"] == 12 * 60
    assert calc2["des_espera_faz"] == 2 * 60 + 30

    carga = {r["viagem"]: r for r in conn.execute(
        "SELECT viagem, carga_bordo FROM carga_bordo WHERE condicao = 'discharging'")}
    assert carga["APN24001"]["carga_bordo"] == 54
    assert carga["APN24002"]["carga_bordo"] == 1439
    bunker = conn.execute(
        "SELECT vlsfo, mgo FROM abastecimento ab JOIN escala e ON e.id = ab.escala_id "
        " WHERE e.codigo_porto = 'ICOARACI'").fetchone()
    assert (bunker["vlsfo"], bunker["mgo"]) == (394.65, 55.574)

    # importar de novo sem --apagar-existente recusa; com, substitui
    with pytest.raises(RuntimeError):
        imp.importar(conn, por_navio, apagar=False)
    de_novo = imp.importar(conn, por_navio, apagar=True)
    assert de_novo["apagado"]["viagem"] == 3
    assert conn.execute("SELECT COUNT(*) FROM viagem").fetchone()[0] == 3


def test_importacao_apaga_a_viagem_de_teste_que_o_sistema_abriu(conn, planilha):
    from app import viagens as servico
    teste, _ = servico.abrir_viagem(conn, 2, por="teste")
    assert teste
    por_navio, _ = hp.ler(planilha)
    imp.importar(conn, por_navio, apagar=True)
    numeros = [r[0] for r in conn.execute("SELECT numero FROM viagem WHERE navio_id = 2 ORDER BY id")]
    assert numeros == ["APN24001", "APN24002", "APN24003"]
