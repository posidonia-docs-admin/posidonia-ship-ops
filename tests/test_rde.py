# -*- coding: utf-8 -*-
"""O boletim do meio-dia (RDE): o formulario GFS1201 preenchido pelo sistema.

Tres garantias: (1) o arquivo gerado E o modelo da Alcoa, com as celulas
certas preenchidas e tudo o mais intacto; (2) o que o sistema sabe chega
pre-preenchido; (3) nada do RDE e gravado no banco.
"""
import io
import json
import re
import urllib.parse
from contextlib import closing
from datetime import date, datetime, timedelta, timezone

import openpyxl
import pytest

from app import db, rde
from tests.test_web import cliente, entrar, escala_de, _lancar  # noqa: F401 (fixtures)


# ---------------------------------------------------------------------------
# 1. Preencher o modelo
# ---------------------------------------------------------------------------

def _abrir(conteudo):
    return openpyxl.load_workbook(io.BytesIO(conteudo))


def test_o_modelo_esta_no_app_e_vem_em_branco():
    assert rde.MODELO.exists()
    ws = openpyxl.load_workbook(rde.MODELO)["RDE"]
    for campo, (celula, _) in rde.CELULAS.items():
        assert ws[celula].value is None, (campo, celula, ws[celula].value)
    assert len(ws.merged_cells.ranges) == 204
    assert len(ws._images) == 1          # o logo


def test_preenche_as_celulas_certas_e_preserva_o_resto():
    conteudo = rde.preencher({
        "embarcacao": "AMAZON PATHFINDER", "imo": 9991109, "data": "2026-10-04",
        "numero_viagem": "APT26018", "posidonia_nr": "APT26018", "cliente_nr": "-",
        "comandante": "LUIZ MESQUITA DA SILVA", "nacionalidade": "BRASILEIRA",
        "localizacao": "TERMINAL DA ALUMAR", "lat": "02°40,67'S", "long": "044°21,63'W",
        "calado_proa": "4", "calado_popa": "7,2", "proximo_porto": "BRJUR", "eta": "09/10/2026 PM",
        "dist_navegar": "987", "horas_on_hire": "24", "tripulantes": "22", "reps_cliente": "0",
        "reps_armador": "0", "passageiros": "0",
        "diario_hora_1": "04:20", "diario_desc_1": "ARRIVAL FAZENDINHA",
        "meteo_vento_12": "9", "meteo_baro_18": "1009",
        "est_vlsfo_anterior": "477.87", "est_vlsfo_total": "473,34",
        "est_agua_anterior": "150", "est_agua_recebido": "80", "est_agua_total": "223",
        "maq_mcp_diarias": "4", "maq_mcp_operando": "9193", "maq_caldeira_operando": "5129",
        "observacoes": "NAVIO ATRACADO AO TERMINAL DA ALUMAR.",
    })
    wb = _abrir(conteudo)
    assert wb.sheetnames == ["Gráfico1", "RDE"]
    ws = wb["RDE"]
    assert ws["G12"].value == "AMAZON PATHFINDER" and ws["P12"].value == 9991109
    assert ws["W12"].value == datetime(2026, 10, 4) and ws["W12"].number_format == "mm-dd-yy"
    assert ws["AB12"].value == "APT26018" and ws["AH15"].value == "APT26018" and ws["AH13"].value == "-"
    assert ws["G14"].value == "LUIZ MESQUITA DA SILVA" and ws["T14"].value == "BRASILEIRA"
    assert ws["H18"].value == "TERMINAL DA ALUMAR" and ws["X18"].value == "02°40,67'S"
    assert ws["AK18"].value == 4 and ws["AK19"].value == 7.2          # virgula aceita
    assert ws["X20"].value == "BRJUR" and ws["AH20"].value == "09/10/2026 PM"
    assert ws["P21"].value == 987 and ws["S22"].value == 24 and ws["J23"].value == 22
    assert ws["AM23"].value == "=SUM(J23,Q23,X23,AF23)"
    assert ws["C27"].value == "04:20" and ws["I27"].value == "ARRIVAL FAZENDINHA" and ws["I28"].value is None
    assert ws["AK27"].value == 9 and ws["AM33"].value == 1009 and ws["AG27"].value is None
    assert ws["H37"].value == 477.87 and ws["T37"].value == 473.34 and ws["L37"].value == "=H37+P37-T37"
    assert ws["H39"].value == 150 and ws["P39"].value == 80 and ws["T39"].value == 223
    assert ws["AD37"].value == 4 and ws["AI37"].value == 9193 and ws["AI41"].value == 5129
    assert ws["C46"].value == "NAVIO ATRACADO AO TERMINAL DA ALUMAR."
    # o que faz o arquivo ser o formulario da Alcoa, e nao uma tabela qualquer
    assert len(ws.merged_cells.ranges) == 204 and len(ws._images) == 1
    assert ws["C12"].value.startswith("Embarcação") and ws["H18"].fill.fgColor.rgb == "FFEAEAEA"


def test_numero_invalido_ou_sem_data_e_recusado():
    with pytest.raises(rde.CampoInvalido, match="calado_proa"):
        rde.preencher({"data": "2026-10-04", "calado_proa": "quatro"})
    with pytest.raises(rde.CampoInvalido, match="data"):
        rde.preencher({"calado_proa": "4"})


def test_campo_vazio_deixa_a_celula_em_branco():
    ws = _abrir(rde.preencher({"data": "2026-10-04", "tripulantes": ""}))["RDE"]
    assert ws["J23"].value is None


# ---------------------------------------------------------------------------
# 2. O e-mail
# ---------------------------------------------------------------------------

def test_o_email_e_o_texto_combinado_palavra_por_palavra():
    e = rde.email("AMAZON PATHFINDER", date(2026, 10, 4))
    assert e["para"] == ["rde@posidoniashipping.com", "ops.dry@posidoniashipping.com"]
    assert e["corpo"] == ("Prezados, boa tarde!\n\nSegue anexo o boletim do meio-dia do navio "
                          "Amazon Pathfinder, referente ao dia 04/10/2026. – 12:00 HORAS.")
    assert e["arquivo"] == "GFS1201 - Relatório Diário da Embarcação RDE 04.10.2026.xlsx"
    assert e["mailto"].startswith("mailto:rde@posidoniashipping.com,ops.dry@posidoniashipping.com?subject=")
    corpo = urllib.parse.parse_qs(e["mailto"].split("?", 1)[1])["body"][0]
    assert corpo == e["corpo"]


# ---------------------------------------------------------------------------
# 3. Pela tela
# ---------------------------------------------------------------------------

def _agora_menos(horas):
    return (datetime.now(timezone(timedelta(hours=-3))) - timedelta(hours=horas)).strftime("%Y-%m-%dT%H:%M")


def test_a_tela_do_comandante_traz_o_botao_e_o_formulario_pre_preenchido(cliente):
    entrar(cliente)
    html = cliente.get("/navio").text
    assert "data-abrir-rde" in html and 'id="modal-rde"' in html and 'id="form-rde"' in html
    assert "/static/rde.js" in html
    # o que o sistema sabe, antes de qualquer marco
    assert 'name="numero_viagem"' in html and 'value="APT' in html
    assert 'name="data"' in html and 'value="{}"'.format(rde.hoje_a_bordo().isoformat()) in html
    assert ">9991109<" in html and ">AMAZON PATHFINDER<" in html
    assert 'name="proximo_porto"' in html and 'value="BRJUR"' in html   # viagem nova: Juruti e o proximo


def test_o_pre_preenchimento_acompanha_a_viagem(cliente):
    entrar(cliente)
    cliente.get("/navio")                       # abre a viagem do piloto
    alumar, fazendinha, juruti = escala_de(cliente, 10), escala_de(cliente, 20), escala_de(cliente, 30)
    # Sailing de Alumar ontem, Arrival em Fazendinha ha 3 horas (dentro das 24h)
    cliente.post("/api/marco", json={
        "escala_id": alumar, "tipo": "sailing", "hora_local": _agora_menos(30), "offset": "-03:00",
        "nome_responsavel": "Cmt.", "id_cliente": "rde-1", "rob_vlsfo": 480, "rob_mgo": 70, "fw": 150})
    cliente.post("/api/marco", json={
        "escala_id": fazendinha, "tipo": "arrival", "hora_local": _agora_menos(3), "offset": "-03:00",
        "nome_responsavel": "Cmt.", "id_cliente": "rde-2", "rob_vlsfo": 473.34, "rob_mgo": 67.47, "fw": 140})
    html = cliente.get("/navio").text
    assert 'name="est_vlsfo_total" value="473.34"' in html and 'name="est_vlsfo_anterior" value="480' in html
    assert 'name="est_agua_total" value="140' in html
    assert 'name="proximo_porto" value="BRJUR"' in html
    assert 'name="localizacao" value="EM FAZENDINHA"' in html
    # o diario: so o marco das ultimas 24h
    assert 'name="diario_desc_1" value="ARRIVAL FAZENDINHA"' in html
    assert 'name="diario_desc_2" value=""' in html
    assert escala_de(cliente, 30) == juruti


def test_gerar_devolve_o_xlsx_e_o_email_sem_gravar_nada(cliente):
    entrar(cliente)
    with closing(db.conectar()) as conn:
        antes = {t: conn.execute("SELECT COUNT(*) FROM {}".format(t)).fetchone()[0]
                 for t in ("evento", "escala", "viagem", "abastecimento", "movimento_carga")}
    r = cliente.post("/navio/rde", json={
        "data": "2026-10-04", "numero_viagem": "APT26018", "comandante": "LUIZ", "tripulantes": "22",
        "embarcacao": "NAVIO FALSO", "imo": "1"})     # navio e IMO vem do cadastro, nao do formulario
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "RDE 04.10.2026.xlsx" in r.headers["content-disposition"]
    email = json.loads(urllib.parse.unquote(r.headers["x-rde-email"]))
    assert email["arquivo"].endswith("RDE 04.10.2026.xlsx") and "Amazon Pathfinder" in email["corpo"]
    ws = _abrir(r.content)["RDE"]
    assert ws["G12"].value == "AMAZON PATHFINDER" and ws["P12"].value == 9991109
    assert ws["G14"].value == "LUIZ" and ws["J23"].value == 22
    with closing(db.conectar()) as conn:
        depois = {t: conn.execute("SELECT COUNT(*) FROM {}".format(t)).fetchone()[0]
                  for t in antes}
    assert depois == antes


def test_dado_invalido_devolve_422_com_o_campo(cliente):
    entrar(cliente)
    r = cliente.post("/navio/rde", json={"data": "2026-10-04", "calado_proa": "x"})
    assert r.status_code == 422 and "calado_proa" in r.json()["erros"][0]


def test_so_a_conta_do_navio_gera_o_rde(cliente):
    entrar(cliente, "vlo")
    assert cliente.post("/navio/rde", json={"data": "2026-10-04"}).status_code == 403


def test_o_rde_nao_reconhece_sof_nem_grava_no_banco():
    """Decisao do Vinicius: nada de SOF, e o RDE nao vira tabela."""
    codigo = rde.MODELO.parent.parent.joinpath("rde.py").read_text(encoding="utf-8")
    assert not re.search(r"INSERT|CREATE TABLE|UPDATE ", codigo)
    assert "sof" not in codigo.lower().replace("isoformat", "")
