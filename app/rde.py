# -*- coding: utf-8 -*-
"""RDE — o Relatório Diário da Embarcação (formulário GFS1201 da Alcoa).

Todo dia, ao meio-dia, o comandante manda por e-mail a planilha GFS1201
preenchida. Este módulo faz três coisas, e só três:

  1. diz o que o sistema JÁ SABE para o dia (navio, IMO, viagem, onde o navio
     está, próximo porto, os ROB dos últimos marcos, os marcos das últimas
     24 horas) — é o que vem pré-preenchido no formulário;
  2. preenche o PRÓPRIO arquivo-modelo da Alcoa (`app/modelos/rde.xlsx`), célula
     por célula, preservando mesclagens, logo e fórmulas;
  3. monta o texto do e-mail que acompanha o anexo.

Nada do que o comandante digita aqui vai para o banco — decisão do Vinícius
em 05/10/2026: o RDE existe para cumprir o envio, não para ser analisado. O que
precisa ser lembrado de um dia para o outro (calado, horas de máquina, tripu-
lação) fica no aparelho do comandante, pelo rde.js.
"""
from __future__ import annotations

import io
import urllib.parse
from datetime import date, datetime, timedelta, timezone

import openpyxl

from . import config, dominio

MODELO = config.RAIZ / "app" / "modelos" / "rde.xlsx"
ABA = "RDE"

DESTINATARIOS = ("rde@posidoniashipping.com", "ops.dry@posidoniashipping.com")
HORA_BOLETIM = "12:00"
FUSO_BORDO = timezone(timedelta(hours=-3))      # toda a rota fica em UTC-3

# As linhas das tabelas do formulário, na ordem em que aparecem na planilha.
DIARIO_LINHAS = 7                               # linhas 27 a 33
METEO_HORAS = ("00", "06", "12", "18")          # colunas AG, AI, AK, AM
METEO = (
    ("vento",   "Vento (nós)"),
    ("direcao", "Vento (direção)"),
    ("ondas",   "Ondas (m)"),
    ("mar",     "Força do mar"),
    ("temp",    "Temperatura (ºC)"),
    ("visib",   "Visibilidade (mn)"),
    ("baro",    "Barômetro"),
)
ESTOQUES = (("vlsfo", "VLSFO"), ("mgo", "MGO"), ("agua", "ÁGUA"))
MAQUINAS = (("mcp", "MCP"), ("mca1", "MCA 01"), ("mca2", "MCA 02"),
            ("mca3", "MCA 03"), ("caldeira", "CALDEIRA"))

# Campo do formulário -> célula (canto superior esquerdo da mesclagem) e tipo.
# "texto" grava como está; "numero" aceita vírgula e grava float; "data" grava
# datetime (a célula tem formato de data no modelo).
CELULAS = {
    "embarcacao":     ("G12", "texto"),
    "imo":            ("P12", "numero"),
    "data":           ("W12", "data"),
    "numero_viagem":  ("AB12", "texto"),
    "cliente_nr":     ("AH13", "texto"),
    "comandante":     ("G14", "texto"),
    "nacionalidade":  ("T14", "texto"),
    "posidonia_nr":   ("AH15", "texto"),
    "localizacao":    ("H18", "texto"),
    "lat":            ("X18", "texto"),
    "long":           ("X19", "texto"),
    "calado_proa":    ("AK18", "numero"),
    "calado_popa":    ("AK19", "numero"),
    "proteina":       ("J20", "texto"),
    "proximo_porto":  ("X20", "texto"),
    "eta":            ("AH20", "texto"),
    "dist_navegar":   ("P21", "numero"),
    "dist_coberta":   ("AE21", "numero"),
    "velocidade":     ("AM21", "numero"),
    "rumo":           ("H22", "numero"),
    "horas_on_hire":  ("S22", "numero"),
    "rancho_dias":    ("AC22", "numero"),
    "residuos":       ("AM22", "numero"),
    "tripulantes":    ("J23", "numero"),
    "reps_cliente":   ("Q23", "numero"),
    "reps_armador":   ("X23", "numero"),
    "passageiros":    ("AF23", "numero"),
    "observacoes":    ("C46", "texto"),
}
for _i in range(1, DIARIO_LINHAS + 1):
    _linha = 26 + _i
    CELULAS["diario_hora_{}".format(_i)] = ("C{}".format(_linha), "texto")
    CELULAS["diario_codigo_{}".format(_i)] = ("F{}".format(_linha), "texto")
    CELULAS["diario_desc_{}".format(_i)] = ("I{}".format(_linha), "texto")
for _i, (_chave, _) in enumerate(METEO):
    for _hora, _col in zip(METEO_HORAS, ("AG", "AI", "AK", "AM")):
        CELULAS["meteo_{}_{}".format(_chave, _hora)] = ("{}{}".format(_col, 27 + _i), "numero")
for _i, (_chave, _) in enumerate(ESTOQUES):
    _linha = 37 + _i
    CELULAS["est_{}_anterior".format(_chave)] = ("H{}".format(_linha), "numero")
    CELULAS["est_{}_recebido".format(_chave)] = ("P{}".format(_linha), "numero")
    CELULAS["est_{}_total".format(_chave)] = ("T{}".format(_linha), "numero")
for _i, (_chave, _) in enumerate(MAQUINAS):
    _linha = 37 + _i
    CELULAS["maq_{}_diarias".format(_chave)] = ("AD{}".format(_linha), "numero")
    CELULAS["maq_{}_operando".format(_chave)] = ("AI{}".format(_linha), "numero")

# Campos que ficam na planilha como fórmula, nunca como número digitado.
FORMULAS = {
    "AM23": "=SUM(J23,Q23,X23,AF23)",           # total embarcado
    "L37": "=H37+P37-T37",                      # consumido = anterior + recebido - ROB
    "L38": "=H38+P38-T38",
    "L39": "=H39+P39-T39",
}


class CampoInvalido(ValueError):
    pass


# ---------------------------------------------------------------------------
# 1. O que o sistema já sabe
# ---------------------------------------------------------------------------

def hoje_a_bordo() -> date:
    return datetime.now(FUSO_BORDO).date()


def _marcos_cronologicos(blocos):
    """(bloco, marco) de todo marco lançado, na ordem da viagem."""
    return [(b, m) for b in blocos for m in b["marcos"] if m["lancado"]]


def _proximo_porto(blocos):
    """A próxima parada em que o navio opera (carrega, descarrega, abastece).

    Passagem e espera de maré não contam: no RDE de referência o navio em
    Alumar anuncia Juruti como próximo porto, não Fazendinha.
    """
    pares = _marcos_cronologicos(blocos)
    indice = blocos.index(pares[-1][0]) if pares else -1
    # Esteja o navio ainda no porto ou já navegando, o próximo é o que vem
    # depois da parada do último marco.
    for bloco in blocos[indice + 1:]:
        escala = bloco["escala"]
        if escala["tipo_escala"] in ("operacional", "encerramento") or escala["motivo"] == "bunker":
            return escala
    return None


def dados_do_sistema(conn, conta, corrente, situacao) -> dict:
    """Os campos do formulário que o sistema consegue preencher sozinho.

    `corrente` é a viagem aberta montada por `_contexto_navio`; `situacao` é a
    frase de onde o navio está (frota.pernada_atual). Tudo é sugestão: o
    comandante pode mudar qualquer campo antes de gerar o arquivo.
    """
    navio = conn.execute("SELECT nome_oficial, imo FROM navio WHERE id = ?",
                         (conta["navio_id"],)).fetchone()
    campos = {
        "embarcacao": navio["nome_oficial"],
        "imo": navio["imo"],
        "data": hoje_a_bordo().isoformat(),
        "nacionalidade": "BRASILEIRA",
        "cliente_nr": "-",
        "horas_on_hire": 24,
    }
    if corrente is None:
        return {"campos": campos, "diario": []}

    blocos = corrente["blocos"]
    campos["numero_viagem"] = corrente["viagem"]["numero"]
    campos["posidonia_nr"] = corrente["viagem"]["numero"]
    if situacao and situacao.get("titulo"):
        campos["localizacao"] = situacao["titulo"].upper()

    proximo = _proximo_porto(blocos)
    if proximo is not None:
        locode = conn.execute("SELECT un_locode FROM porto WHERE codigo = ?",
                              (proximo["codigo_porto"],)).fetchone()
        campos["proximo_porto"] = (locode["un_locode"] if locode and locode["un_locode"]
                                   else proximo["porto_nome"].upper())

    # ROB: o último marco é o "Total ROB"; o marco anterior serve de "Qtd
    # anterior" quando o aparelho não lembra o boletim de ontem.
    pares = _marcos_cronologicos(blocos)
    ultimo = pares[-1][1]["lancado"] if pares else None
    penultimo = pares[-2][1]["lancado"] if len(pares) > 1 else None
    for chave, coluna in (("vlsfo", "rob_vlsfo"), ("mgo", "rob_mgo"), ("agua", "fw")):
        if ultimo is not None and ultimo[coluna] is not None:
            campos["est_{}_total".format(chave)] = ultimo[coluna]
        if penultimo is not None and penultimo[coluna] is not None:
            campos["est_{}_anterior".format(chave)] = penultimo[coluna]

    # Recebido nas últimas 24 horas: os abastecimentos gravados na viagem.
    limite = (datetime.now(timezone.utc) - timedelta(hours=24)).strftime("%Y-%m-%dT%H:%M:%S")
    recebido = conn.execute(
        "SELECT COALESCE(SUM(a.vlsfo), 0) AS vlsfo, COALESCE(SUM(a.mgo), 0) AS mgo "
        "  FROM abastecimento a JOIN escala e ON e.id = a.escala_id "
        " WHERE e.viagem_id = ? AND a.registrado_em >= ?",
        (corrente["viagem"]["id"], limite)).fetchone()
    if recebido and recebido["vlsfo"]:
        campos["est_vlsfo_recebido"] = recebido["vlsfo"]
    if recebido and recebido["mgo"]:
        campos["est_mgo_recebido"] = recebido["mgo"]

    # O diário de operações: os marcos das últimas 24 horas, hora local e
    # descrição ("ARRIVAL FAZENDINHA"). O comandante completa o resto.
    diario = []
    for bloco, marco in pares:
        if marco["lancado"]["hora_utc"] >= limite:
            diario.append({
                "hora": marco["lancado"]["hora_local"][11:16],
                "codigo": "",
                "desc": "{} {}".format(dominio.ROTULO_CURTO[marco["tipo"]],
                                       bloco["escala"]["porto_nome"]).upper(),
            })
    return {"campos": campos, "diario": diario[:DIARIO_LINHAS]}


# ---------------------------------------------------------------------------
# 2. Preencher o modelo
# ---------------------------------------------------------------------------

def _numero(bruto):
    texto = str(bruto).strip().replace(",", ".")
    try:
        valor = float(texto)
    except ValueError:
        raise CampoInvalido("'{}' não é um número.".format(bruto)) from None
    return int(valor) if valor.is_integer() else valor


def _data(bruto):
    try:
        return datetime.strptime(str(bruto).strip(), "%Y-%m-%d")
    except ValueError:
        raise CampoInvalido("Data inválida: '{}'.".format(bruto)) from None


def normalizar(bruto: dict) -> dict:
    """Converte o que veio do formulário (tudo texto) nos valores da planilha.

    Campo vazio vira None — a célula fica em branco, como no modelo. Erro em
    qualquer campo numérico ou na data para tudo, com o nome do campo.
    """
    limpo = {}
    erros = []
    for campo, (_, tipo) in CELULAS.items():
        valor = bruto.get(campo)
        if valor is None or str(valor).strip() == "":
            limpo[campo] = None
            continue
        try:
            if tipo == "numero":
                limpo[campo] = _numero(valor)
            elif tipo == "data":
                limpo[campo] = _data(valor)
            else:
                limpo[campo] = str(valor).strip()
        except CampoInvalido as exc:
            erros.append("{}: {}".format(campo, exc))
    if limpo.get("data") is None:
        erros.append("data: informe a data do boletim.")
    if erros:
        raise CampoInvalido(" ".join(erros))
    return limpo


def preencher(dados: dict) -> bytes:
    """O arquivo GFS1201 com os campos preenchidos, pronto para anexar."""
    limpo = normalizar(dados)
    wb = openpyxl.load_workbook(MODELO)
    ws = wb[ABA]
    for campo, (celula, _) in CELULAS.items():
        ws[celula] = limpo[campo]
    for celula, formula in FORMULAS.items():
        ws[celula] = formula
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


# ---------------------------------------------------------------------------
# 3. O e-mail
# ---------------------------------------------------------------------------

def _nome_bonito(nome_oficial: str) -> str:
    """'AMAZON PATHFINDER' -> 'Amazon Pathfinder'."""
    return " ".join(p.capitalize() for p in (nome_oficial or "").split())


def nome_arquivo(dia: date) -> str:
    return "GFS1201 - Relatório Diário da Embarcação RDE {}.xlsx".format(dia.strftime("%d.%m.%Y"))


def email(nome_oficial: str, dia: date) -> dict:
    """Destinatários, assunto, corpo e o link mailto: que abre o e-mail pronto.

    O texto é o que o Vinícius passou em 05/10/2026, palavra por palavra; só
    o navio e a data mudam. O anexo o comandante junta à mão — nenhum
    navegador anexa arquivo por conta própria.
    """
    navio = _nome_bonito(nome_oficial)
    data_br = dia.strftime("%d/%m/%Y")
    assunto = "Boletim do meio-dia – {} – {}".format(navio, data_br)
    corpo = ("Prezados, boa tarde!\n\n"
             "Segue anexo o boletim do meio-dia do navio {}, referente ao dia {}. "
             "– {} HORAS.".format(navio, data_br, HORA_BOLETIM))
    mailto = "mailto:{}?{}".format(
        ",".join(DESTINATARIOS),
        urllib.parse.urlencode({"subject": assunto, "body": corpo}, quote_via=urllib.parse.quote))
    return {"para": list(DESTINATARIOS), "assunto": assunto, "corpo": corpo,
            "mailto": mailto, "arquivo": nome_arquivo(dia)}
